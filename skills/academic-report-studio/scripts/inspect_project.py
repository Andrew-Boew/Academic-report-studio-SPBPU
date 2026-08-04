#!/usr/bin/env python3
"""Create a privacy-conscious inventory of files used for an academic report."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import mimetypes
import os
from pathlib import Path
from typing import Any


IGNORED_DIRS = {
    ".git",
    ".idea",
    ".vscode",
    "__pycache__",
    "node_modules",
    "output",
    "renders",
    "tmp",
}
TEXT_EXTENSIONS = {
    ".c", ".cc", ".cpp", ".css", ".h", ".hpp", ".html", ".ipynb",
    ".java", ".js", ".jsx", ".md", ".m", ".py", ".r", ".rs", ".sh",
    ".sql", ".swift", ".tex", ".ts", ".tsx", ".txt", ".yaml", ".yml",
}
MAX_HASH_BYTES = 250 * 1024 * 1024
PROFILE_FILENAMES = {".report-profile.json", ".report-subjects.json"}


def sha256(path: Path) -> str | None:
    if path.stat().st_size > MAX_HASH_BYTES:
        return None
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_text(path: Path, limit: int = 256_000) -> tuple[str, str]:
    raw = path.read_bytes()[:limit]
    for encoding in ("utf-8-sig", "utf-8", "cp1251", "latin-1"):
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace"), "utf-8-replace"


def inspect_csv(path: Path) -> dict[str, Any]:
    text, encoding = read_text(path)
    lines = text.splitlines()
    sample = "\n".join(lines[:30])
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    rows = list(csv.reader(lines[:6], delimiter=delimiter))
    header = rows[0] if rows else []
    return {
        "kind": "table",
        "encoding": encoding,
        "delimiter": delimiter,
        "estimated_rows": max(0, len(lines) - 1),
        "columns": len(header),
        "headers": header,
    }


def inspect_xlsx(path: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        return {"kind": "workbook", "warning": f"openpyxl unavailable: {exc}"}
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheets = []
    for sheet in workbook.worksheets:
        header = []
        if sheet.max_row:
            header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        sheets.append(
            {
                "name": sheet.title,
                "rows": sheet.max_row,
                "columns": sheet.max_column,
                "headers": header,
                "state": sheet.sheet_state,
            }
        )
    workbook.close()
    return {"kind": "workbook", "sheets": sheets}


def inspect_json(path: Path) -> dict[str, Any]:
    text, encoding = read_text(path, limit=5_000_000)
    try:
        value = json.loads(text)
    except json.JSONDecodeError as exc:
        return {"kind": "json", "encoding": encoding, "error": str(exc)}
    result: dict[str, Any] = {"kind": "json", "encoding": encoding}
    if isinstance(value, list):
        result["root_type"] = "array"
        result["items"] = len(value)
        if value and isinstance(value[0], dict):
            result["keys"] = sorted(map(str, value[0].keys()))
    elif isinstance(value, dict):
        result["root_type"] = "object"
        result["keys"] = sorted(map(str, value.keys()))
    else:
        result["root_type"] = type(value).__name__
    return result


def inspect_docx(path: Path) -> dict[str, Any]:
    try:
        from docx import Document
    except ImportError as exc:
        return {"kind": "docx", "warning": f"python-docx unavailable: {exc}"}
    document = Document(path)
    headings = [
        p.text.strip()
        for p in document.paragraphs
        if p.text.strip() and p.style and p.style.name.startswith("Heading")
    ]
    return {
        "kind": "docx",
        "sections": len(document.sections),
        "paragraphs": sum(bool(p.text.strip()) for p in document.paragraphs),
        "tables": len(document.tables),
        "inline_images": len(document.inline_shapes),
        "headings": headings[:100],
    }


def inspect_pdf(path: Path) -> dict[str, Any]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        return {"kind": "pdf", "warning": f"pypdf unavailable: {exc}"}
    reader = PdfReader(path)
    encrypted = bool(reader.is_encrypted)
    return {"kind": "pdf", "pages": len(reader.pages), "encrypted": encrypted}


def inspect_image(path: Path) -> dict[str, Any]:
    try:
        from PIL import Image
        with Image.open(path) as image:
            return {
                "kind": "image",
                "width_px": image.width,
                "height_px": image.height,
                "mode": image.mode,
                "format": image.format,
            }
    except Exception as exc:  # Pillow reports format-specific errors.
        return {"kind": "image", "warning": str(exc)}


def inspect_text(path: Path) -> dict[str, Any]:
    text, encoding = read_text(path)
    return {
        "kind": "text_or_code",
        "encoding": encoding,
        "lines": text.count("\n") + (1 if text else 0),
        "nonempty_lines": sum(bool(line.strip()) for line in text.splitlines()),
    }


def inspect_file(path: Path, root: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()
    item: dict[str, Any] = {
        "path": str(path.resolve()),
        "relative_path": str(path.relative_to(root)),
        "extension": suffix,
        "mime_type": mimetypes.guess_type(path.name)[0],
        "size_bytes": path.stat().st_size,
        "sha256": sha256(path),
    }
    if suffix in {".csv", ".tsv"}:
        item.update(inspect_csv(path))
    elif suffix in {".xlsx", ".xlsm"}:
        item.update(inspect_xlsx(path))
    elif suffix == ".json":
        item.update(inspect_json(path))
        if path.name == ".report-profile.json":
            item["kind"] = "report_profile"
        elif path.name == ".report-subjects.json":
            item["kind"] = "subject_profiles"
    elif suffix == ".docx":
        item.update(inspect_docx(path))
    elif suffix == ".pdf":
        item.update(inspect_pdf(path))
    elif suffix in {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".webp"}:
        item.update(inspect_image(path))
    elif suffix in TEXT_EXTENSIONS:
        item.update(inspect_text(path))
    else:
        item["kind"] = "other"
    return item


def collect(root: Path) -> list[Path]:
    paths = []
    for current_root, dirnames, filenames in os.walk(root):
        dirnames[:] = [name for name in dirnames if name not in IGNORED_DIRS and not name.startswith(".")]
        for name in filenames:
            if (name.startswith(".") and name not in PROFILE_FILENAMES) or name.startswith("~$"):
                continue
            paths.append(Path(current_root) / name)
    return sorted(paths)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_dir", type=Path)
    parser.add_argument("--out", required=True, type=Path)
    args = parser.parse_args()

    root = args.input_dir.expanduser().resolve()
    if not root.is_dir():
        parser.error(f"Input directory does not exist: {root}")

    items = []
    errors = []
    for path in collect(root):
        try:
            items.append(inspect_file(path, root))
        except Exception as exc:
            errors.append({"path": str(path), "error": f"{type(exc).__name__}: {exc}"})

    report = {
        "root": str(root),
        "file_count": len(items),
        "total_size_bytes": sum(item["size_bytes"] for item in items),
        "files": items,
        "errors": errors,
        "privacy_note": "Headers and structural metadata only; dataset rows are not copied into this inventory.",
    }
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(args.out.resolve()), "files": len(items), "errors": len(errors)}, ensure_ascii=False))
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
