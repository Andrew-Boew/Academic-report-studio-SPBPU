#!/usr/bin/env python3
"""Validate structural and formatting invariants of a generated academic DOCX."""

from __future__ import annotations

import argparse
import json
import re
import zipfile
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


PLACEHOLDER = re.compile(r"\{\{|\}\}|\b(?:TODO|TBD|FIXME)\b|\[вставить[^\]]*\]", re.I)
OBJECT_NUMBER = r"(?:[А-ЯЁ]\.\d+|\d+(?:\.\d+)*)"
TABLE_CAPTION = re.compile(rf"^Таблица\s+({OBJECT_NUMBER})\s+—\s+\S.*$", re.I)
TABLE_CONTINUATION = re.compile(rf"^(Продолжение|Окончание)\s+таблицы\s+({OBJECT_NUMBER})$", re.I)
FIGURE_CAPTION = re.compile(rf"^Рисунок\s+({OBJECT_NUMBER})\s+—\s+\S.*$", re.I)
LISTING_CAPTION = re.compile(rf"^Листинг\s+({OBJECT_NUMBER})\s+—\s+\S.*$", re.I)
MANUAL_HEADING_NUMBER = re.compile(r"^\s*\d+(?:\.\d+)*[.)]?\s+")


def near(value: float, expected: float, tolerance: float = 0.03) -> bool:
    return abs(value - expected) <= tolerance


def style_font_name(style) -> str | None:
    return style.font.name


def xml_text(element) -> str:
    return "".join(node.text or "" for node in element.findall(".//" + qn("w:t"))).strip()


def is_layout_table(table_element) -> bool:
    """Identify the borderless cover-details table, which is not a numbered data table."""
    borders = table_element.find("./" + qn("w:tblPr") + "/" + qn("w:tblBorders"))
    if borders is None or len(borders) == 0:
        return False
    values = [border.get(qn("w:val"), "") for border in borders]
    return bool(values) and all(value in {"nil", "none"} for value in values)


def is_code_fragment(table_element) -> bool:
    descriptor = table_element.find("./" + qn("w:tblPr") + "/" + qn("w:tblCaption"))
    return descriptor is not None and descriptor.get(qn("w:val")) == "code-fragment"


def is_figure_placeholder(table_element) -> bool:
    descriptor = table_element.find("./" + qn("w:tblPr") + "/" + qn("w:tblCaption"))
    return descriptor is not None and descriptor.get(qn("w:val")) == "figure-placeholder"


def effective_alignment(paragraph):
    return paragraph.alignment if paragraph.alignment is not None else paragraph.style.paragraph_format.alignment


def warn_numeric_gaps(numbers: list[str], label: str, warnings: list[str]) -> None:
    if numbers and all(number.isdigit() for number in numbers):
        values = [int(number) for number in numbers]
        expected = list(range(1, len(values) + 1))
        if values != expected:
            warnings.append(f"Нарушена последовательная сквозная нумерация {label}: {values}")


def inspect(path: Path) -> dict[str, Any]:
    errors: list[str] = []
    warnings: list[str] = []

    if path.suffix.lower() != ".docx":
        errors.append("Итоговый файл должен иметь расширение .docx")
        return {"path": str(path), "errors": errors, "warnings": warnings, "passed": False}
    if not path.is_file():
        errors.append("Файл не найден")
        return {"path": str(path), "errors": errors, "warnings": warnings, "passed": False}
    if not zipfile.is_zipfile(path):
        errors.append("Файл не является корректным контейнером DOCX/ZIP")
        return {"path": str(path), "errors": errors, "warnings": warnings, "passed": False}

    try:
        document = Document(path)
    except Exception as exc:
        errors.append(f"DOCX не открывается: {type(exc).__name__}: {exc}")
        return {"path": str(path), "errors": errors, "warnings": warnings, "passed": False}

    if not document.sections:
        errors.append("В документе отсутствуют секции")
    for index, section in enumerate(document.sections, 1):
        prefix = f"Секция {index}"
        if section.orientation != WD_ORIENT.PORTRAIT:
            errors.append(f"{prefix}: требуется строго книжная ориентация страницы")
        expected_width = 21.0
        expected_height = 29.7
        if not near(section.page_width.cm, expected_width) or not near(section.page_height.cm, expected_height):
            errors.append(
                f"{prefix}: ожидается формат A4, получено "
                f"{section.page_width.cm:.2f} × {section.page_height.cm:.2f} см"
            )
        checks = (
            ("левое поле", section.left_margin.cm, 3.0),
            ("правое поле", section.right_margin.cm, 1.5),
            ("верхнее поле", section.top_margin.cm, 2.0),
            ("нижнее поле", section.bottom_margin.cm, 2.0),
        )
        for label, actual, expected in checks:
            if not near(actual, expected):
                errors.append(f"{prefix}: {label} {actual:.2f} см вместо {expected:.2f} см")

    first_section = document.sections[0]
    if not first_section.different_first_page_header_footer:
        errors.append("На титульном листе не отключён показ номера страницы")
    page_numbering = first_section._sectPr.find(qn("w:pgNumType"))
    page_start = page_numbering.get(qn("w:start")) if page_numbering is not None else None
    if page_start != "0":
        errors.append("На титульном листе номер должен быть скрыт, а первая страница содержания должна иметь номер 1")
    footer_paragraphs = first_section.footer.paragraphs
    if not footer_paragraphs or effective_alignment(footer_paragraphs[0]) != WD_ALIGN_PARAGRAPH.RIGHT:
        errors.append("Номер страницы должен располагаться справа внизу")

    normal = document.styles["Normal"]
    font_name = style_font_name(normal)
    if font_name != "Times New Roman":
        errors.append(f"Стиль Normal: шрифт {font_name!r} вместо Times New Roman")
    if normal.font.size is None or not near(normal.font.size.pt, 14.0, 0.1):
        errors.append("Стиль Normal: размер должен быть 14 pt")
    if normal.paragraph_format.first_line_indent is None or not near(normal.paragraph_format.first_line_indent.cm, 1.25):
        errors.append("Стиль Normal: абзацный отступ должен быть 1,25 см")
    if normal.paragraph_format.alignment != WD_ALIGN_PARAGRAPH.JUSTIFY:
        errors.append("Стиль Normal: требуется выравнивание по ширине")
    spacing = normal.paragraph_format.line_spacing
    if not isinstance(spacing, float) or not near(spacing, 1.5, 0.01):
        errors.append("Стиль Normal: межстрочный интервал должен быть 1,5")

    expected_heading_sizes = {"Heading 1": 16.0, "Heading 2": 14.0, "Heading 3": 14.0}
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[name]
        expected_size = expected_heading_sizes[name]
        if style.font.name != "Times New Roman" or style.font.size is None or not near(style.font.size.pt, expected_size, 0.1):
            errors.append(f"Стиль {name}: требуется Times New Roman {expected_size:g} pt")
        if style.paragraph_format.alignment != WD_ALIGN_PARAGRAPH.LEFT:
            errors.append(f"Стиль {name}: нумерованные заголовки должны быть выровнены слева")
        first_indent = style.paragraph_format.first_line_indent
        first_indent_cm = first_indent.cm if first_indent is not None else 0.0
        left_indent = style.paragraph_format.left_indent
        left_indent_cm = left_indent.cm if left_indent is not None else 0.0
        if not near(first_indent_cm, 0.0) or not near(left_indent_cm, 0.0):
            errors.append(
                f"Стиль {name}: прямые отступы должны быть 0 см; общий отступ 1,25 см задаёт нумерация Word"
            )
        if style.paragraph_format.space_before is None or not near(style.paragraph_format.space_before.pt, 0.0, 0.1):
            errors.append(f"Стиль {name}: интервал перед должен быть 0 pt")
        if style.paragraph_format.space_after is None or not near(style.paragraph_format.space_after.pt, 0.0, 0.1):
            errors.append(f"Стиль {name}: интервал после должен быть 0 pt")
    if not document.styles["Heading 3"].font.italic:
        errors.append("Стиль Heading 3 должен быть полужирным курсивом")

    heading_num_ids: list[str] = []
    for level, name in enumerate(("Heading 1", "Heading 2", "Heading 3")):
        ppr = document.styles[name]._element.get_or_add_pPr()
        num_pr = ppr.find(qn("w:numPr"))
        ilvl = num_pr.find(qn("w:ilvl")) if num_pr is not None else None
        num_id = num_pr.find(qn("w:numId")) if num_pr is not None else None
        if ilvl is None or ilvl.get(qn("w:val")) != str(level) or num_id is None:
            errors.append(f"Стиль {name} не связан с уровнем {level + 1} многоуровневой нумерации Word")
            continue
        heading_num_ids.append(num_id.get(qn("w:val"), ""))
    if heading_num_ids and len(set(heading_num_ids)) != 1:
        errors.append("Heading 1–3 должны использовать одну многоуровневую нумерацию Word")
    if heading_num_ids and len(set(heading_num_ids)) == 1:
        numbering = document.part.numbering_part.element
        num_node = next(
            (node for node in numbering.findall(qn("w:num")) if node.get(qn("w:numId")) == heading_num_ids[0]),
            None,
        )
        abstract_ref = num_node.find(qn("w:abstractNumId")) if num_node is not None else None
        abstract_id = abstract_ref.get(qn("w:val")) if abstract_ref is not None else None
        abstract = next(
            (
                node for node in numbering.findall(qn("w:abstractNum"))
                if node.get(qn("w:abstractNumId")) == abstract_id
            ),
            None,
        )
        for level in range(3):
            lvl = next(
                (
                    node for node in abstract.findall(qn("w:lvl"))
                    if node.get(qn("w:ilvl")) == str(level)
                ),
                None,
            ) if abstract is not None else None
            indent = lvl.find(qn("w:pPr") + "/" + qn("w:ind")) if lvl is not None else None
            if (
                indent is None
                or indent.get(qn("w:left")) != "1417"
                or indent.get(qn("w:hanging")) != "708"
            ):
                errors.append(
                    f"Уровень {level + 1} нумерации заголовков должен начинать номер на 1,25 см и текст на 2,50 см"
                )

    if "Report Structural Heading" not in [style.name for style in document.styles]:
        errors.append("Не найден стиль ненумерованных структурных заголовков")
    else:
        structural_style = document.styles["Report Structural Heading"]
        structural_num_pr = structural_style._element.get_or_add_pPr().find(qn("w:numPr"))
        if structural_num_pr is not None:
            errors.append("Структурные заголовки не должны наследовать нумерацию глав")

    for name in ("Report Cover", "Report Caption", "Report Table Caption", "Report Equation"):
        if name not in [style.name for style in document.styles]:
            errors.append(f"Не найден стиль {name}")
            continue
        style = document.styles[name]
        if style.font.name != "Times New Roman" or style.font.size is None or not near(style.font.size.pt, 14.0, 0.1):
            if name != "Report Equation":
                errors.append(f"Стиль {name}: требуется Times New Roman 14 pt")
        if name in {"Report Caption", "Report Table Caption", "Report Equation"}:
            line_spacing = style.paragraph_format.line_spacing
            before = style.paragraph_format.space_before
            after = style.paragraph_format.space_after
            if not isinstance(line_spacing, float) or not near(line_spacing, 1.5, 0.01):
                errors.append(f"Стиль {name}: межстрочный интервал должен быть 1,5")
            if before is None or not near(before.pt, 0.0, 0.1):
                errors.append(f"Стиль {name}: интервал перед должен быть 0 pt")
            if after is None or not near(after.pt, 0.0, 0.1):
                errors.append(f"Стиль {name}: интервал после должен быть 0 pt")

    first_page_text = "\n".join(p.text for p in document.paragraphs[:35])
    report_type = "course" if re.search(r"КУРСОВАЯ\s+РАБОТА", first_page_text, re.I) else "lab"
    if "политехнический университет" not in first_page_text.lower():
        errors.append("Не найден обязательный титульный блок университета")
    if not re.search(r"ЛАБОРАТОРНАЯ\s+РАБОТА|КУРСОВАЯ\s+РАБОТА", first_page_text, re.I):
        errors.append("На титульном листе не указан тип учебной работы")
    work_label_paragraphs = [
        p for p in document.paragraphs
        if re.search(r"ЛАБОРАТОРНАЯ\s+РАБОТА|КУРСОВАЯ\s+РАБОТА", p.text, re.I)
    ]
    if work_label_paragraphs:
        runs = [run for run in work_label_paragraphs[0].runs if run.text.strip()]
        if not runs or any(run.font.size is None or not near(run.font.size.pt, 14.0, 0.1) for run in runs):
            errors.append("Тип работы на титульном листе должен иметь размер 14 pt")
        if any(run.font.bold is not True for run in runs):
            errors.append("Тип работы на титульном листе должен быть полужирным")
    signature_paragraphs = [p for p in document.paragraphs[:35] if p.text.strip() == "<подпись>"]
    if len(signature_paragraphs) != 2:
        errors.append("На титульном листе должны быть две строки <подпись>: для студента и проверяющего")
    for paragraph in signature_paragraphs:
        if effective_alignment(paragraph) != WD_ALIGN_PARAGRAPH.CENTER:
            errors.append("Строка <подпись> должна быть выровнена по центру")
        runs = [run for run in paragraph.runs if run.text.strip()]
        if any(run.font.size is None or not near(run.font.size.pt, 12.0, 0.1) or run.font.italic is not True for run in runs):
            errors.append("Строка <подпись> должна быть Times New Roman 12 pt курсивом")

    all_text = "\n".join(p.text for p in document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                all_text += "\n" + cell.text
    if PLACEHOLDER.search(all_text):
        errors.append("В документе остались служебные заполнители")

    with zipfile.ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml")
        settings_xml = archive.read("word/settings.xml")
        footer_xml = b"".join(
            archive.read(name) for name in archive.namelist() if name.startswith("word/footer") and name.endswith(".xml")
        )
    has_toc = b"TOC " in document_xml
    if not has_toc:
        warnings.append("Поле содержания TOC не найдено")
    if b"updateFields" not in settings_xml:
        errors.append("Не установлен автоматический запрос на обновление полей Word")
    if b"PAGE " not in footer_xml:
        errors.append("Поле номера страницы PAGE не найдено в нижнем колонтитуле")
    if "Обновите содержание в Word".encode("utf-8") in document_xml:
        warnings.append("После открытия в Word нужно обновить содержание (Ctrl+A, F9)")

    toc_expected_indents = {"TOC 1": 0.0, "TOC 2": 1.25, "TOC 3": 2.5}
    for style_name, expected_indent in toc_expected_indents.items():
        if style_name not in [style.name for style in document.styles]:
            continue
        indent = document.styles[style_name].paragraph_format.left_indent
        actual_indent = indent.cm if indent is not None else 0.0
        if not near(actual_indent, expected_indent):
            errors.append(
                f"Стиль {style_name}: левый отступ {actual_indent:.2f} см вместо {expected_indent:.2f} см"
            )
        line_spacing = document.styles[style_name].paragraph_format.line_spacing
        if line_spacing is None or not near(float(line_spacing), 1.5, 0.05):
            errors.append(
                f"Стиль {style_name}: межстрочный интервал содержания должен быть полуторным (1,5)"
            )

    toc_paragraphs = [
        paragraph for paragraph in document.paragraphs
        if paragraph.style and paragraph.style.name in toc_expected_indents
    ]
    for paragraph in toc_paragraphs:
        style_name = paragraph.style.name
        expected_indent = toc_expected_indents[style_name]
        direct_indent = paragraph.paragraph_format.left_indent
        style_indent = document.styles[style_name].paragraph_format.left_indent
        effective_indent = (
            direct_indent.cm
            if direct_indent is not None
            else (style_indent.cm if style_indent is not None else 0.0)
        )
        if not near(effective_indent, expected_indent):
            errors.append(
                f"Строка {style_name}: левый отступ {effective_indent:.2f} см вместо {expected_indent:.2f} см"
            )
        direct_first = paragraph.paragraph_format.first_line_indent
        effective_first = direct_first.cm if direct_first is not None else 0.0
        if not near(effective_first, 0.0):
            errors.append(f"Строка {style_name}: отступ первой строки содержания должен быть 0 см")
    update_fields = document.settings._element.find(qn("w:updateFields"))
    update_value = update_fields.get(qn("w:val"), "true").lower() if update_fields is not None else ""
    if toc_paragraphs and update_value not in {"false", "0", "off"}:
        errors.append(
            "После материализации содержания автоматическое updateFields должно быть отключено, чтобы Word не показывал нули"
        )
    if has_toc and not toc_paragraphs and update_value not in {"true", "1", "on"}:
        errors.append("Черновое содержание должно запрашивать обновление полей перед материализацией")
    for paragraph in toc_paragraphs:
        tabs = paragraph._p.xpath("./w:pPr/w:tabs/w:tab[@w:val='right'][@w:leader='dot']")
        if not tabs:
            errors.append(f"Строка содержания не имеет правой табуляции с точечным заполнителем: {paragraph.text!r}")
            continue
        if tabs[0].get(qn("w:pos")) != "9354":
            errors.append("Правый табулятор содержания должен находиться внутри рабочей ширины 16,5 см")
        text_nodes = paragraph._p.xpath(".//w:t/text()")
        if not text_nodes or not text_nodes[-1].strip().isdigit():
            errors.append(f"В строке содержания отсутствует номер страницы: {paragraph.text!r}")
        elif int(text_nodes[-1].strip()) <= 0:
            errors.append(f"В строке содержания указан нулевой номер страницы: {paragraph.text!r}")

    bookmarks = {
        node.get(qn("w:name"))
        for node in document.element.xpath(".//w:bookmarkStart")
        if node.get(qn("w:name"))
    }
    internal_links = document.element.xpath(".//w:hyperlink[@w:anchor]")
    if has_toc and not internal_links:
        warnings.append("Внутренние гиперссылки содержания ещё не материализованы")
    for link in internal_links:
        anchor = link.get(qn("w:anchor"), "")
        if anchor not in bookmarks:
            errors.append(f"Внутренняя гиперссылка ведёт на отсутствующую закладку: {anchor}")

    for link in document.element.xpath(".//w:hyperlink"):
        for run in link.xpath(".//w:r"):
            if not xml_text(run):
                continue
            rpr = run.find("./" + qn("w:rPr"))
            color = rpr.find("./" + qn("w:color")) if rpr is not None else None
            underline = rpr.find("./" + qn("w:u")) if rpr is not None else None
            size = rpr.find("./" + qn("w:sz")) if rpr is not None else None
            fonts = rpr.find("./" + qn("w:rFonts")) if rpr is not None else None
            if color is None or color.get(qn("w:val"), "").upper() != "000000":
                errors.append("Гиперссылка должна оставаться чёрной")
                break
            if underline is None or underline.get(qn("w:val"), "") != "none":
                errors.append("Гиперссылка не должна иметь подчёркивание")
                break
            if size is None or size.get(qn("w:val")) != "28":
                errors.append("Текст гиперссылки должен иметь размер 14 pt")
                break
            if fonts is None or fonts.get(qn("w:ascii")) != "Times New Roman" or fonts.get(qn("w:hAnsi")) != "Times New Roman":
                errors.append("Текст гиперссылки должен иметь шрифт Times New Roman")
                break

    heading_style_names = {"Heading 1", "Heading 2", "Heading 3", "Report Structural Heading"}
    heading_texts = [
        p.text.strip() for p in document.paragraphs
        if p.style and p.style.name in heading_style_names
    ]
    if not heading_texts:
        errors.append("Не найдены настоящие стили заголовков Word")
    for paragraph in document.paragraphs:
        if not paragraph.style or paragraph.style.name not in heading_style_names:
            continue
        value = paragraph.text.strip()
        if paragraph.style.name.startswith("Heading "):
            if MANUAL_HEADING_NUMBER.match(value):
                errors.append(f"Номер главы набран вручную вместо функции нумерации Word: {value!r}")
            direct_left = paragraph.paragraph_format.left_indent
            direct_first = paragraph.paragraph_format.first_line_indent
            left_cm = direct_left.cm if direct_left is not None else 0.0
            first_cm = direct_first.cm if direct_first is not None else 0.0
            if not near(left_cm, 0.0) or not near(first_cm, 1.25):
                errors.append(
                    f"{paragraph.style.name}: номер заголовка должен начинаться на 1,25 см "
                    f"(прямой левый отступ 0 см, первая строка 1,25 см)"
                )
        words = re.findall(r"[А-ЯЁа-яё]{3,}", value)
        if words and len(words) >= 2 and all(word.isupper() for word in words):
            errors.append(f"Заголовок должен быть в регистре предложения, а не прописными: {value!r}")

    appendix_indices = [i for i, text in enumerate(heading_texts) if text.upper().startswith("ПРИЛОЖЕНИЕ")]
    source_indices = [i for i, text in enumerate(heading_texts) if "ИСТОЧНИК" in text.upper() or "ЛИТЕРАТУР" in text.upper()]
    if appendix_indices and source_indices and min(appendix_indices) < max(source_indices):
        errors.append("Приложения должны располагаться после списка источников")

    bibliography_paragraphs = [
        paragraph for paragraph in document.paragraphs
        if paragraph.style and paragraph.style.name == "Report Bibliography"
    ]
    if report_type == "course" and not bibliography_paragraphs:
        errors.append("Курсовая работа должна содержать список проверенных реальных источников")
    if report_type == "lab" and bibliography_paragraphs:
        errors.append("Лабораторная работа не должна содержать список использованных источников")
    seen_external_links = 0
    for paragraph in document.paragraphs:
        external_links = paragraph._p.xpath(".//w:hyperlink[@r:id]")
        if external_links and (not paragraph.style or paragraph.style.name != "Report Bibliography"):
            errors.append("Внешние гиперссылки разрешены только в списке использованных источников")
        if paragraph.style and paragraph.style.name == "Report Bibliography":
            if len(external_links) != 1:
                errors.append("Каждый источник должен содержать ровно одну проверенную внешнюю гиперссылку")
        for link in external_links:
            seen_external_links += 1
            relation_id = link.get(qn("r:id"), "")
            relation = paragraph.part.rels.get(relation_id)
            target = relation.target_ref if relation is not None else ""
            parsed = urlparse(target)
            host = (parsed.hostname or "").casefold()
            if parsed.scheme != "https" or not host or host in {"localhost", "127.0.0.1", "::1"}:
                errors.append(f"Источник должен вести на публичный HTTPS-адрес: {target!r}")
    all_external_links = document.element.xpath(".//w:hyperlink[@r:id]")
    if seen_external_links != len(all_external_links):
        errors.append("Обнаружена внешняя гиперссылка вне обычного элемента списка источников")
    if bibliography_paragraphs and seen_external_links != len(bibliography_paragraphs):
        errors.append("Число гиперссылок не совпадает с числом источников")
    if report_type == "lab" and seen_external_links:
        errors.append("В лабораторной работе запрещены внешние гиперссылки")

    narrative_citations: list[int] = []
    for paragraph in document.paragraphs:
        if not paragraph.style or paragraph.style.name not in {"Normal", "List Bullet", "List Number"}:
            continue
        narrative_citations.extend(int(value) for value in re.findall(r"\[(\d+)\]", paragraph.text))
    if report_type == "lab" and narrative_citations:
        errors.append("В лабораторной работе не должно быть ссылок на литературу вида [N]")
    if report_type == "course":
        for number in narrative_citations:
            if number < 1 or number > len(bibliography_paragraphs):
                errors.append(f"Ссылка [{number}] не соответствует элементу списка источников")
        for link in document.element.xpath(".//w:hyperlink[starts-with(@w:anchor, '_ReportBib')]"):
            anchor = link.get(qn("w:anchor"), "")
            if anchor not in bookmarks:
                errors.append(f"Ссылка на источник ведёт на отсутствующую закладку: {anchor}")

    ordinary = [
        p for p in document.paragraphs
        if p.text.strip() and p.style and p.style.name == "Normal"
    ]
    bad_indents = 0
    for paragraph in ordinary:
        direct = paragraph.paragraph_format.first_line_indent
        if (
            direct is not None
            and not near(direct.cm, 1.25)
            and not (near(direct.cm, 0.0) and effective_alignment(paragraph) in {WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT})
        ):
            bad_indents += 1
    if bad_indents:
        warnings.append(f"Обычных абзацев с прямым отступом не 1,25 см: {bad_indents}")
    left_aligned_body = [
        p for p in ordinary
        if effective_alignment(p) == WD_ALIGN_PARAGRAPH.LEFT
    ]
    if left_aligned_body:
        errors.append(
            "Обычный текст должен быть выровнен по ширине; абзацев с выравниванием влево: "
            f"{len(left_aligned_body)}"
        )

    figure_numbers: list[str] = []
    table_numbers: list[str] = []
    listing_numbers: list[str] = []
    continuation_numbers: list[str] = []
    equation_numbers: list[str] = []
    for paragraph in document.paragraphs:
        text_value = paragraph.text.strip()
        if not text_value:
            continue
        figure_match = FIGURE_CAPTION.fullmatch(text_value)
        table_match = TABLE_CAPTION.fullmatch(text_value)
        continuation_match = TABLE_CONTINUATION.fullmatch(text_value)
        listing_match = LISTING_CAPTION.fullmatch(text_value)
        if text_value.lower().startswith("рисунок"):
            if not figure_match:
                errors.append(f"Некорректная подпись рисунка: {text_value!r}")
            else:
                figure_numbers.append(figure_match.group(1).upper())
            if text_value.endswith("."):
                errors.append(f"Точка в конце подписи рисунка: {text_value!r}")
            if effective_alignment(paragraph) != WD_ALIGN_PARAGRAPH.CENTER:
                errors.append(f"Подпись рисунка должна быть по центру: {text_value!r}")
            spacing = paragraph.paragraph_format
            effective_before = spacing.space_before or paragraph.style.paragraph_format.space_before
            effective_after = spacing.space_after or paragraph.style.paragraph_format.space_after
            effective_line = spacing.line_spacing or paragraph.style.paragraph_format.line_spacing
            if effective_before is None or not near(effective_before.pt, 0.0, 0.1):
                errors.append(f"Подпись рисунка должна иметь интервал перед 0 pt: {text_value!r}")
            if effective_after is None or not near(effective_after.pt, 0.0, 0.1):
                errors.append(f"Подпись рисунка должна иметь интервал после 0 pt: {text_value!r}")
            if not isinstance(effective_line, float) or not near(effective_line, 1.5, 0.01):
                errors.append(f"Подпись рисунка должна иметь межстрочный интервал 1,5: {text_value!r}")
        elif text_value.lower().startswith("таблица"):
            if not table_match:
                errors.append(f"Некорректная подпись таблицы: {text_value!r}")
            else:
                table_numbers.append(table_match.group(1).upper())
            if text_value.endswith("."):
                errors.append(f"Точка в конце подписи таблицы: {text_value!r}")
            if effective_alignment(paragraph) != WD_ALIGN_PARAGRAPH.CENTER:
                errors.append(f"Подпись таблицы должна быть по центру: {text_value!r}")
        elif re.match(r"^(?:продолжение|окончание)\s+таблицы", text_value, re.I):
            if not continuation_match:
                errors.append(f"Некорректная подпись переноса таблицы: {text_value!r}")
            else:
                continuation_numbers.append(continuation_match.group(2).upper())
            if effective_alignment(paragraph) != WD_ALIGN_PARAGRAPH.CENTER:
                errors.append(f"Подпись продолжения таблицы должна быть по центру: {text_value!r}")
        elif text_value.lower().startswith("листинг"):
            if not listing_match:
                errors.append(f"Некорректная подпись листинга: {text_value!r}")
            else:
                listing_numbers.append(listing_match.group(1).upper())
            if text_value.endswith("."):
                errors.append(f"Точка в конце подписи листинга: {text_value!r}")
            if effective_alignment(paragraph) != WD_ALIGN_PARAGRAPH.CENTER:
                errors.append(f"Подпись листинга должна быть по центру: {text_value!r}")
        if paragraph.style and paragraph.style.name.startswith("Report Equation"):
            math_objects = paragraph._p.xpath("./m:oMath")
            if not math_objects:
                errors.append("Формула оформлена обычным текстом, а не Microsoft Word Equation (OMML)")
            elif len(math_objects) != 1:
                errors.append("Один формульный абзац должен содержать ровно один объект Microsoft Word Equation")
            spacing = paragraph.paragraph_format
            effective_before = spacing.space_before or paragraph.style.paragraph_format.space_before
            effective_after = spacing.space_after or paragraph.style.paragraph_format.space_after
            effective_line = spacing.line_spacing or paragraph.style.paragraph_format.line_spacing
            if effective_before is None or not near(effective_before.pt, 0.0, 0.1):
                errors.append("Абзац формулы должен иметь интервал перед 0 pt")
            if effective_after is None or not near(effective_after.pt, 0.0, 0.1):
                errors.append("Абзац формулы должен иметь интервал после 0 pt")
            if not isinstance(effective_line, float) or not near(effective_line, 1.5, 0.01):
                errors.append("Абзац формулы должен иметь межстрочный интервал 1,5")
            math_sizes = []
            math_text = "".join(node.text or "" for node in paragraph._p.xpath(".//m:t"))
            for math_run in paragraph._p.xpath(".//m:r"):
                size = math_run.find("./" + qn("w:rPr") + "/" + qn("w:sz"))
                if size is not None:
                    math_sizes.append(int(size.get(qn("w:val"), "0")))
            longest_number = max((len(value) for value in re.findall(r"\d+", math_text)), default=0)
            expected_size = 28
            if longest_number > 60:
                errors.append(
                    "Целое число длиннее 60 цифр должно быть представлено именованными "
                    "десятичными блоками в отдельных формулах 14 pt"
                )
            if not math_sizes or any(value != expected_size for value in math_sizes):
                errors.append(
                    "Размер Microsoft Word Equation не соответствует правилу длинного числа: "
                    f"ожидается {expected_size / 2:g} pt"
                )
            paragraph_size = paragraph._p.find(
                "./" + qn("w:pPr") + "/" + qn("w:rPr") + "/" + qn("w:sz")
            )
            if paragraph_size is None or paragraph_size.get(qn("w:val")) != str(expected_size):
                errors.append(
                    "Размер формулы должен быть продублирован в свойствах формульного абзаца "
                    "для одинакового рендера в Microsoft Word и LibreOffice"
                )
            for matrix in paragraph._p.xpath(".//m:m"):
                rows = matrix.findall("./" + qn("m:mr"))
                if len(rows) <= 1:
                    continue
                equation_rows = []
                for row in rows:
                    row_text = "".join(node.text or "" for node in row.iter(qn("m:t")))
                    cells = row.findall("./" + qn("m:e"))
                    equation_rows.append(len(cells) == 1 and "=" in row_text)
                if all(equation_rows):
                    errors.append(
                        "Самостоятельные равенства или параметры нельзя объединять в одну "
                        "многострочную матрицу; требуется отдельный центрированный абзац для каждой строки"
                    )
            equation_match = re.search(rf"\(({OBJECT_NUMBER})\)\s*$", text_value, re.I)
            if equation_match:
                tabs = paragraph._p.xpath("./w:pPr/w:tabs/w:tab")
                tab_values = {tab.get(qn("w:val")) for tab in tabs}
                if not {"center", "right"}.issubset(tab_values):
                    errors.append("Нумерованная формула должна иметь центральную и правую табуляцию")
            elif effective_alignment(paragraph) != WD_ALIGN_PARAGRAPH.CENTER:
                errors.append("Ненумерованная формула или параметр должны быть по центру")
            if equation_match:
                equation_numbers.append(equation_match.group(1).upper())

    duplicate_figures = sorted({number for number in figure_numbers if figure_numbers.count(number) > 1})
    duplicate_tables = sorted({number for number in table_numbers if table_numbers.count(number) > 1})
    duplicate_equations = sorted({number for number in equation_numbers if equation_numbers.count(number) > 1})
    duplicate_listings = sorted({number for number in listing_numbers if listing_numbers.count(number) > 1})
    if duplicate_figures:
        errors.append(f"Повторяются номера рисунков: {', '.join(duplicate_figures)}")
    if duplicate_tables:
        errors.append(f"Повторяются номера таблиц: {', '.join(duplicate_tables)}")
    if duplicate_equations:
        errors.append(f"Повторяются номера формул: {', '.join(duplicate_equations)}")
    if duplicate_listings:
        errors.append(f"Повторяются номера листингов: {', '.join(duplicate_listings)}")
    for number in continuation_numbers:
        if number not in table_numbers:
            errors.append(f"Продолжение ссылается на отсутствующую таблицу {number}")
    warn_numeric_gaps(figure_numbers, "рисунков", warnings)
    warn_numeric_gaps(table_numbers, "таблиц", warnings)
    warn_numeric_gaps(equation_numbers, "формул", warnings)
    warn_numeric_gaps(listing_numbers, "листингов", warnings)

    narrative_parts = []
    for paragraph in document.paragraphs:
        value = paragraph.text.strip()
        if not value or FIGURE_CAPTION.fullmatch(value) or TABLE_CAPTION.fullmatch(value) or TABLE_CONTINUATION.fullmatch(value) or LISTING_CAPTION.fullmatch(value):
            continue
        if paragraph.style and paragraph.style.name == "Report Equation":
            continue
        narrative_parts.append(value)
    narrative_text = "\n".join(narrative_parts)
    for number in figure_numbers:
        if not re.search(rf"\bрисунк\w*\s+{re.escape(number)}\b", narrative_text, re.I):
            errors.append(f"В тексте отсутствует ссылка на рисунок {number}")
    for number in table_numbers:
        if not re.search(rf"\bтабл(?:иц\w*|\.)\s+{re.escape(number)}\b", narrative_text, re.I):
            errors.append(f"В тексте отсутствует ссылка на таблицу {number}")
    for number in listing_numbers:
        if not re.search(rf"\bлистинг\w*\s+{re.escape(number)}\b", narrative_text, re.I):
            errors.append(f"В тексте отсутствует ссылка на листинг {number}")
    for number in equation_numbers:
        if not re.search(rf"(?:формул\w*|уравнен\w*)\s*\({re.escape(number)}\)", narrative_text, re.I):
            warnings.append(f"В тексте не найдена явная ссылка на формулу ({number})")

    body_items = list(document.element.body.iterchildren())
    seen_first_page_break = False
    content_tables = 0
    for item_index, item in enumerate(body_items):
        if item.tag == qn("w:p"):
            page_break = any(
                node.get(qn("w:type")) == "page"
                for node in item.findall(".//" + qn("w:br"))
            )
            if page_break:
                seen_first_page_break = True
            has_drawing = item.find(".//" + qn("w:drawing")) is not None or item.find(".//" + qn("w:pict")) is not None
            if has_drawing and seen_first_page_break:
                next_text = ""
                for following in body_items[item_index + 1:]:
                    if following.tag == qn("w:p") and not xml_text(following):
                        continue
                    next_text = xml_text(following) if following.tag == qn("w:p") else "<table>"
                    break
                if not FIGURE_CAPTION.fullmatch(next_text):
                    errors.append("Подпись рисунка должна находиться непосредственно под изображением")
        elif item.tag == qn("w:tbl") and not is_layout_table(item):
            if is_figure_placeholder(item):
                next_text = ""
                for following in body_items[item_index + 1:]:
                    if following.tag == qn("w:p") and not xml_text(following):
                        continue
                    next_text = xml_text(following) if following.tag == qn("w:p") else "<table>"
                    break
                if not FIGURE_CAPTION.fullmatch(next_text):
                    errors.append("Подпись места под рисунок должна находиться непосредственно под рамкой")
                continue
            if is_code_fragment(item):
                previous_text = ""
                for previous in reversed(body_items[:item_index]):
                    if previous.tag == qn("w:p") and not xml_text(previous):
                        continue
                    previous_text = xml_text(previous) if previous.tag == qn("w:p") else "<table>"
                    break
                if not LISTING_CAPTION.fullmatch(previous_text):
                    errors.append("Рамка короткого кода должна иметь подпись `Листинг N — Название` непосредственно сверху")
                rows = item.findall("./" + qn("w:tr"))
                cells = item.findall("./" + qn("w:tr") + "/" + qn("w:tc"))
                if len(rows) != 1 or len(cells) != 1:
                    errors.append("Короткий фрагмент кода должен находиться в одноячеечной рамке")
                for run in item.findall(".//" + qn("w:r")):
                    if not xml_text(run):
                        continue
                    rpr = run.find("./" + qn("w:rPr"))
                    fonts = rpr.find("./" + qn("w:rFonts")) if rpr is not None else None
                    size = rpr.find("./" + qn("w:sz")) if rpr is not None else None
                    font_name = fonts.get(qn("w:ascii")) if fonts is not None else None
                    half_points = int(size.get(qn("w:val"), "0")) if size is not None else 0
                    if font_name != "Courier New" or half_points not in {22, 24}:
                        errors.append("Короткий код в рамке должен быть Courier New 11 или 12 pt")
                        break
                continue
            content_tables += 1
            previous_text = ""
            for previous in reversed(body_items[:item_index]):
                if previous.tag == qn("w:p") and not xml_text(previous):
                    continue
                previous_text = xml_text(previous) if previous.tag == qn("w:p") else "<table>"
                break
            if not (TABLE_CAPTION.fullmatch(previous_text) or TABLE_CONTINUATION.fullmatch(previous_text)):
                errors.append(f"Таблица {content_tables}: подпись должна находиться непосредственно сверху")

            table_rows = item.findall("./" + qn("w:tr"))
            if not table_rows:
                errors.append(f"Таблица {content_tables}: отсутствуют строки")
                continue
            header = table_rows[0].find("./" + qn("w:trPr") + "/" + qn("w:tblHeader"))
            if header is None or header.get(qn("w:val"), "true") not in {"true", "1", "on"}:
                errors.append(f"Таблица {content_tables}: строка заголовка не отмечена для повтора")
            if len(table_rows) > 1:
                for header_paragraph in table_rows[0].findall(".//" + qn("w:p")):
                    keep_next = header_paragraph.find("./" + qn("w:pPr") + "/" + qn("w:keepNext"))
                    if keep_next is None:
                        errors.append(
                            f"Таблица {content_tables}: шапка должна удерживаться минимум "
                            "с первой строкой данных"
                        )
                        break
            table_style = item.find("./" + qn("w:tblPr") + "/" + qn("w:tblStyle"))
            if table_style is None or table_style.get(qn("w:val")) not in {"TableGrid", "Table Grid"}:
                errors.append(f"Таблица {content_tables}: требуются тонкие внешние и внутренние границы")
            small_runs = []
            for row_index, row in enumerate(table_rows, 1):
                for cell_index, cell in enumerate(row.findall("./" + qn("w:tc")), 1):
                    no_wrap = cell.find("./" + qn("w:tcPr") + "/" + qn("w:noWrap"))
                    fit_text = cell.find("./" + qn("w:tcPr") + "/" + qn("w:tcFitText"))
                    if no_wrap is not None or fit_text is not None:
                        errors.append(
                            f"Таблица {content_tables}: w:noWrap/w:tcFitText запрещены; "
                            "длинное значение должно переноситься естественно"
                        )
                    cell_text = xml_text(cell).strip().replace("−", "-")
                    if re.fullmatch(r"[+-]?\d+(?:[.,]\d+)?", cell_text):
                        for cell_paragraph in cell.findall("./" + qn("w:p")):
                            alignment = cell_paragraph.find(
                                "./" + qn("w:pPr") + "/" + qn("w:jc")
                            )
                            if alignment is None or alignment.get(qn("w:val")) != "center":
                                errors.append(
                                    f"Таблица {content_tables}: числовая ячейка "
                                    f"{row_index}:{cell_index} должна быть по центру"
                                )
                                break
                    for run in cell.findall(".//" + qn("w:r")):
                        if not xml_text(run):
                            continue
                        rpr = run.find("./" + qn("w:rPr"))
                        size = rpr.find("./" + qn("w:sz")) if rpr is not None else None
                        fonts = rpr.find("./" + qn("w:rFonts")) if rpr is not None else None
                        if size is not None:
                            half_points = int(size.get(qn("w:val"), "0"))
                            font_name = fonts.get(qn("w:ascii")) if fonts is not None else None
                            if half_points < 28:
                                small_runs.append(f"{row_index}:{cell_index}={half_points / 2:g}pt")
                        if row_index == 1:
                            bold = run.find("./" + qn("w:rPr") + "/" + qn("w:b"))
                            if bold is None or bold.get(qn("w:val"), "true") not in {"true", "1", "on"}:
                                errors.append(f"Таблица {content_tables}: текст шапки должен быть полужирным")
                                break
            if small_runs:
                errors.append(
                    f"Таблица {content_tables}: табличный текст меньше обязательных 14 pt: "
                    + ", ".join(small_runs[:10])
                )
            shaded_cells = []
            for row_index, row in enumerate(table_rows, 1):
                for cell_index, cell in enumerate(row.findall("./" + qn("w:tc")), 1):
                    shading = cell.find("./" + qn("w:tcPr") + "/" + qn("w:shd"))
                    if shading is None:
                        continue
                    fill = shading.get(qn("w:fill"), "").upper()
                    if fill not in {"", "AUTO", "FFFFFF", "CLEAR"}:
                        shaded_cells.append(f"{row_index}:{cell_index}={fill}")
            if shaded_cells:
                errors.append(
                    f"Таблица {content_tables}: заливка ячеек запрещена локальным профилем: "
                    + ", ".join(shaded_cells[:10])
                )
            split_rows = []
            for row_index, row in enumerate(table_rows, 1):
                cant_split = row.find("./" + qn("w:trPr") + "/" + qn("w:cantSplit"))
                if cant_split is None or cant_split.get(qn("w:val"), "true") not in {"true", "1", "on"}:
                    split_rows.append(row_index)
            if split_rows:
                errors.append(
                    f"Таблица {content_tables}: Word может разорвать строки между страницами: "
                    + ", ".join(map(str, split_rows[:10]))
                )

    return {
        "path": str(path),
        "report_type": report_type,
        "size_bytes": path.stat().st_size,
        "sections": len(document.sections),
        "paragraphs": sum(bool(p.text.strip()) for p in document.paragraphs),
        "tables": len(document.tables),
        "content_tables": content_tables,
        "figure_numbers": figure_numbers,
        "table_numbers": table_numbers,
        "listing_numbers": listing_numbers,
        "equation_numbers": equation_numbers,
        "inline_images": len(document.inline_shapes),
        "bibliography_sources": len(bibliography_paragraphs),
        "external_source_links": seen_external_links,
        "errors": errors,
        "warnings": warnings,
        "passed": not errors,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx", type=Path)
    parser.add_argument("--json", dest="json_path", type=Path)
    args = parser.parse_args()
    result = inspect(args.docx.expanduser().resolve())
    payload = json.dumps(result, ensure_ascii=False, indent=2)
    if args.json_path:
        args.json_path.parent.mkdir(parents=True, exist_ok=True)
        args.json_path.write_text(payload, encoding="utf-8")
    print(payload)
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
