#!/usr/bin/env python3
"""Build a Russian academic report DOCX from a constrained JSON specification."""

from __future__ import annotations

import argparse
import copy
import datetime as dt
import functools
import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import urlparse

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor
from lxml import etree


MAX_CONTENT_WIDTH_CM = 16.5
PLACEHOLDER_PATTERN = re.compile(r"\{\{|\}\}|\b(?:TODO|TBD|FIXME)\b|\[вставить[^\]]*\]", re.I)
MANUAL_HEADING_NUMBER = re.compile(r"^\s*\d+(?:\.\d+)*[.)]?\s+")
CITATION_PATTERN = re.compile(r"\[(\d+)\]")
SOURCE_KINDS = {
    "article",
    "book",
    "conference_paper",
    "dataset",
    "official_documentation",
    "official_web_resource",
    "standard",
}
OPENING_PAGE_HEADINGS = {
    "введение",
    "цель работы",
    "задачи",
}
DEFAULT_NEW_PAGE_HEADINGS = {
    "ход работы",
    "дополнительные задания",
    "дополнительное задание",
    "выводы",
    "заключение",
}


class SpecError(ValueError):
    pass


def set_run_font(run, name: str, size: float, *, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), name)
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn(f"w:{key}"), None)


def set_character_spacing(run, points: float) -> None:
    """Set expanded character spacing while keeping searchable text intact."""
    spacing = run._element.get_or_add_rPr().find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        run._element.get_or_add_rPr().append(spacing)
    spacing.set(qn("w:val"), str(round(points * 20)))


def add_external_hyperlink(paragraph, text: str, url: str) -> None:
    """Add a clickable URL that remains black and un-underlined per the local profile."""
    relation_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), "Times New Roman")
    rpr.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "28")
    rpr.append(size)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "none")
    rpr.append(underline)
    run.append(rpr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_internal_hyperlink(paragraph, text: str, anchor: str) -> None:
    """Add a black, un-underlined link to a bookmark in the same document."""
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), "Times New Roman")
    rpr.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "28")
    rpr.append(size)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "none")
    rpr.append(underline)
    run.append(rpr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_report_text(paragraph, text: str) -> None:
    """Write body text and turn course-paper [N] citations into internal links."""
    cursor = 0
    for match in CITATION_PATTERN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, "Times New Roman", 14)
        number = int(match.group(1))
        add_internal_hyperlink(paragraph, match.group(0), f"_ReportBib{number}")
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, "Times New Roman", 14)


def configure_style(style, *, font="Times New Roman", size=14, bold=None, italic=None,
                    align=None, first_indent_cm=None, left_indent_cm=None,
                    line_spacing=1.5, before=0, after=0, keep_next=None) -> None:
    style.font.name = font
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = RGBColor(0, 0, 0)
    rfonts = style._element.get_or_add_rPr().rFonts
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), font)
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(qn(f"w:{key}"), None)
    pf = style.paragraph_format
    pf.alignment = align
    pf.first_line_indent = Cm(first_indent_cm) if first_indent_cm is not None else None
    pf.left_indent = Cm(left_indent_cm) if left_indent_cm is not None else None
    pf.line_spacing = line_spacing
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.keep_with_next = keep_next
    pf.widow_control = True


def ensure_style(document: Document, name: str, style_type=WD_STYLE_TYPE.PARAGRAPH):
    try:
        return document.styles[name]
    except KeyError:
        return document.styles.add_style(name, style_type)


def set_outline_level(style, level: int) -> None:
    ppr = style._element.get_or_add_pPr()
    node = ppr.find(qn("w:outlineLvl"))
    if node is None:
        node = OxmlElement("w:outlineLvl")
        ppr.append(node)
    node.set(qn("w:val"), str(level))


def set_style_numbering(style, *, level: int, num_id: int) -> None:
    ppr = style._element.get_or_add_pPr()
    old = ppr.find(qn("w:numPr"))
    if old is not None:
        ppr.remove(old)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)
    ppr.append(num_pr)


def configure_heading_numbering(document: Document) -> int:
    """Create a real Word multilevel list linked to Heading 1/2/3."""
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId"), "0"))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(node.get(qn("w:numId"), "0"))
        for node in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "multilevel")
    abstract.append(multi)
    for level, label in enumerate(("%1.", "%1.%2.", "%1.%2.%3.")):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "decimal")
        p_style = OxmlElement("w:pStyle")
        p_style.set(qn("w:val"), f"Heading{level + 1}")
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), label)
        suffix = OxmlElement("w:suff")
        suffix.set(qn("w:val"), "space")
        justification = OxmlElement("w:lvlJc")
        justification.set(qn("w:val"), "left")
        ppr = OxmlElement("w:pPr")
        indent = OxmlElement("w:ind")
        # Number position = left - hanging = 709 twips (1.25 cm).
        # Heading text begins at 1417 twips (2.50 cm); Heading 1–3 themselves
        # keep zero direct indents so the positions are never added twice.
        indent.set(qn("w:left"), "1417")
        indent.set(qn("w:hanging"), "708")
        ppr.append(indent)
        for child in (start, num_fmt, p_style, lvl_text, suffix, justification, ppr):
            lvl.append(child)
        abstract.append(lvl)

    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(numbering.index(first_num), abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    reference = OxmlElement("w:abstractNumId")
    reference.set(qn("w:val"), str(abstract_id))
    num.append(reference)
    numbering.append(num)

    for level in range(3):
        set_style_numbering(document.styles[f"Heading {level + 1}"], level=level, num_id=num_id)
    return num_id


def configure_document(document: Document, page_start: int) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)
    section.different_first_page_header_footer = True

    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn("w:start"), str(page_start))

    normal = document.styles["Normal"]
    configure_style(
        normal,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_indent_cm=1.25,
        line_spacing=1.5,
    )

    configure_style(
        document.styles["Heading 1"], size=16, bold=True,
        align=WD_ALIGN_PARAGRAPH.LEFT, first_indent_cm=0, left_indent_cm=0,
        line_spacing=1.5, before=0, after=0, keep_next=True,
    )
    configure_style(
        document.styles["Heading 2"], size=14, bold=True,
        align=WD_ALIGN_PARAGRAPH.LEFT, first_indent_cm=0, left_indent_cm=0,
        line_spacing=1.5, before=0, after=0, keep_next=True,
    )
    configure_style(
        document.styles["Heading 3"], size=14, bold=True, italic=True,
        align=WD_ALIGN_PARAGRAPH.LEFT, first_indent_cm=0, left_indent_cm=0,
        line_spacing=1.5, before=0, after=0, keep_next=True,
    )
    structural_heading = ensure_style(document, "Report Structural Heading")
    configure_style(
        structural_heading, size=16, bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER, first_indent_cm=0,
        line_spacing=1.5, before=0, after=0, keep_next=True,
    )
    set_outline_level(structural_heading, 0)
    configure_heading_numbering(document)

    configure_style(
        ensure_style(document, "Report Caption"), size=14,
        align=WD_ALIGN_PARAGRAPH.CENTER, first_indent_cm=0,
        line_spacing=1.5, before=0, after=0, keep_next=False,
    )
    configure_style(
        ensure_style(document, "Report Table Caption"), size=14,
        align=WD_ALIGN_PARAGRAPH.CENTER, first_indent_cm=0,
        line_spacing=1.5, before=0, after=0, keep_next=True,
    )
    configure_style(
        ensure_style(document, "Report Code Fragment"), font="Courier New", size=12,
        align=WD_ALIGN_PARAGRAPH.LEFT, first_indent_cm=0,
        line_spacing=1.0, before=0, after=0,
    )
    configure_style(
        ensure_style(document, "Report Appendix Code"), font="Courier New", size=10,
        align=WD_ALIGN_PARAGRAPH.LEFT, first_indent_cm=0,
        line_spacing=1.0, before=0, after=0,
    )
    configure_style(
        ensure_style(document, "Report Appendix File Label"), font="Times New Roman", size=12,
        bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent_cm=0,
        line_spacing=1.0, before=6, after=0, keep_next=True,
    )
    configure_style(
        ensure_style(document, "Report Figure Placeholder"), size=12, italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER, first_indent_cm=0,
        line_spacing=1.0, before=0, after=0, keep_next=True,
    )
    configure_style(
        ensure_style(document, "Report Equation"), font="Cambria Math", size=14,
        align=WD_ALIGN_PARAGRAPH.CENTER, first_indent_cm=0,
        line_spacing=1.5, before=0, after=0,
    )
    for half_points in (28,):
        equation_style = ensure_style(document, f"Report Equation {half_points}")
        equation_style.base_style = document.styles["Report Equation"]
        configure_style(
            equation_style, font="Cambria Math", size=half_points / 2,
            align=WD_ALIGN_PARAGRAPH.CENTER, first_indent_cm=0,
            line_spacing=1.5, before=0, after=0,
        )
    configure_style(
        ensure_style(document, "Report Cover"), size=14,
        align=WD_ALIGN_PARAGRAPH.CENTER, first_indent_cm=0,
        line_spacing=1.15, before=0, after=0,
    )
    configure_style(
        ensure_style(document, "Report TOC Field"), size=14,
        align=WD_ALIGN_PARAGRAPH.LEFT, first_indent_cm=0,
        line_spacing=1.5, before=0, after=0,
    )
    configure_style(
        ensure_style(document, "Report Bibliography"), size=14,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent_cm=-0.75,
        left_indent_cm=1.25, line_spacing=1.5,
    )

    for name, left, first in (
        ("List Bullet", 2.5, -0.63),
        ("List Number", 2.5, -0.63),
    ):
        style = document.styles[name]
        configure_style(
            style, size=14, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
            first_indent_cm=first, left_indent_cm=left,
            line_spacing=1.5,
        )

    toc_left_indents_cm = {1: 0.0, 2: 1.25, 3: 2.5}
    for level in range(1, 4):
        name = f"TOC {level}"
        configure_style(
            ensure_style(document, name), size=14,
            align=WD_ALIGN_PARAGRAPH.LEFT,
            first_indent_cm=0,
            left_indent_cm=toc_left_indents_cm[level],
            line_spacing=1.5,
            after=0,
        )

    add_page_number(section.footer.paragraphs[0])
    add_update_fields(document)


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))
    set_run_font(run, "Times New Roman", 14)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.first_line_indent = Cm(0)
    add_field(paragraph, " PAGE ", "1")


def add_update_fields(document: Document) -> None:
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def add_centered_paragraph(document: Document, text: str = "", size: float = 14,
                           bold: bool = False, after: float = 0) -> Any:
    paragraph = document.add_paragraph(style="Report Cover")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_after = Pt(after)
    run = paragraph.add_run(text)
    set_run_font(run, "Times New Roman", size, bold=bold)
    return paragraph


def required(metadata: dict[str, Any], key: str) -> str:
    value = str(metadata.get(key, "")).strip()
    if not value:
        raise SpecError(f"Missing required metadata field: {key}")
    return value


def abbreviate_name(value: str, *, spaced_initials: bool) -> str:
    """Convert a full Russian name to the surname-and-initials title-page form."""
    value = value.strip()
    if "." in value:
        return value
    parts = [part for part in value.split() if part]
    if len(parts) < 2:
        return value
    separator = " " if spaced_initials else ""
    initials = separator.join(f"{part[0].upper()}." for part in parts[1:])
    return f"{parts[0]} {initials}"


def capitalized_label(value: str) -> str:
    value = value.strip()
    return value[:1].upper() + value[1:] if value else value


def add_cover(document: Document, spec: dict[str, Any]) -> None:
    meta = spec["metadata"]
    report_type = spec["report_type"]

    ministry = str(meta.get("ministry", "Министерство науки и высшего образования Российской Федерации")).strip()
    organization_type = str(
        meta.get(
            "organization_type",
            "Федеральное государственное автономное образовательное учреждение высшего образования",
        )
    ).strip()
    institution = str(meta.get("institution", "Санкт-Петербургский политехнический университет Петра Великого")).strip()
    institute = str(meta.get("institute", "Институт компьютерных наук и кибербезопасности")).strip()
    department = str(meta.get("department", "Высшая школа кибербезопасности")).strip()

    if ministry:
        add_centered_paragraph(document, ministry)
    if organization_type:
        add_centered_paragraph(document, organization_type)
    if institution:
        add_centered_paragraph(document, f"«{institution.strip('«»')}»")
    if institute:
        add_centered_paragraph(document, institute)
    if department:
        department_paragraph = add_centered_paragraph(document, department, bold=True)
        department_paragraph.paragraph_format.space_after = Pt(154)

    if report_type == "lab":
        number = str(meta.get("work_number", "")).strip()
        work_label = "Лабораторная работа" + (f" № {number}" if number else "")
    else:
        work_label = "Курсовая работа"
    add_centered_paragraph(document, work_label, size=14, bold=True, after=16)
    work_title = required(meta, "title").strip("«»\"")
    add_centered_paragraph(document, f"«{work_title}»", size=14, bold=True, after=8)
    variant = str(meta.get("variant", "")).strip()
    if variant:
        add_centered_paragraph(document, f"Вариант {variant}", size=14, after=4)
    discipline = required(meta, "discipline").strip("«»\"")
    discipline_paragraph = add_centered_paragraph(document, f"по дисциплине «{discipline}»", size=14)
    discipline_paragraph.paragraph_format.space_after = Pt(86)

    student_role = str(meta.get("student_role", "Выполнил")).strip()
    student_status = str(meta.get("student_status", "студент")).strip()
    student_display_name = str(meta.get("student_display_name", "")).strip() or abbreviate_name(
        required(meta, "student_name"), spaced_initials=False
    )
    reviewer_display_name = str(meta.get("reviewer_display_name", "")).strip() or abbreviate_name(
        required(meta, "reviewer_name"), spaced_initials=True
    )
    reviewer_position = required(meta, "reviewer_position")
    detail_rows = [
        (student_role, "", 0.0),
        (f"{student_status} гр. {required(meta, 'student_group')}", student_display_name, 0.0),
    ]
    for left, right, after in detail_rows:
        paragraph = document.add_paragraph(style="Report Cover")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_after = Pt(after)
        paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(MAX_CONTENT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
        if left:
            run = paragraph.add_run(left)
            set_run_font(run, "Times New Roman", 14)
        if right:
            run = paragraph.add_run(f"\t{right}")
            set_run_font(run, "Times New Roman", 14)

    signature = add_centered_paragraph(document, "<подпись>", size=12, after=66)
    for run in signature.runs:
        run.font.italic = True

    reviewer_rows = [
        (str(meta.get("reviewer_role", "Проверил:")).strip(), "", 0.0),
        (reviewer_position, reviewer_display_name, 0.0),
    ]
    for left, right, after in reviewer_rows:
        paragraph = document.add_paragraph(style="Report Cover")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_after = Pt(after)
        paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(MAX_CONTENT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
        if left:
            set_run_font(paragraph.add_run(left), "Times New Roman", 14)
        if right:
            set_run_font(paragraph.add_run(f"\t{right}"), "Times New Roman", 14)

    signature = add_centered_paragraph(document, "<подпись>", size=12, after=93)
    for run in signature.runs:
        run.font.italic = True

    add_centered_paragraph(document, required(meta, "city"), size=14)
    add_centered_paragraph(document, required(meta, "year"), size=14)


def add_toc(document: Document, levels: int) -> None:
    heading = document.add_paragraph("Содержание", style="Report Structural Heading")
    heading.paragraph_format.page_break_before = True
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.first_line_indent = Cm(0)
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "9")
    heading._p.get_or_add_pPr().append(outline)
    paragraph = document.add_paragraph(style="Report TOC Field")
    paragraph.paragraph_format.first_line_indent = Cm(0)
    add_field(paragraph, f' TOC \\o "1-{levels}" \\h \\z \\u ', "Обновите содержание в Word")


def clean_text(value: Any, *, field: str) -> str:
    text = str(value if value is not None else "").strip()
    if PLACEHOLDER_PATTERN.search(text):
        raise SpecError(f"Placeholder detected in {field}: {text[:80]}")
    return text


def keep_paragraph_lines(paragraph) -> None:
    paragraph.paragraph_format.widow_control = True


def keep_previous_reference(document: Document, *, max_characters: int = 180) -> None:
    """Keep a short reference paragraph with the following caption/object."""
    for paragraph in reversed(document.paragraphs):
        if not paragraph.text.strip():
            continue
        # Keeping a long paragraph with a following figure/table can move the
        # whole paragraph to the next page and leave half of the prior page
        # empty. Only a genuinely short lead-in is kept with the object.
        if len(paragraph.text) <= max_characters:
            paragraph.paragraph_format.keep_with_next = True
        return


def add_heading(document: Document, block: dict[str, Any]) -> None:
    level = int(block.get("level", 1))
    if level not in (1, 2, 3):
        raise SpecError(f"Heading level must be 1, 2, or 3: {level}")
    text = clean_text(block.get("text"), field="heading")
    if MANUAL_HEADING_NUMBER.match(text):
        raise SpecError(
            f"Heading text must not contain a manually typed number; Word numbers headings automatically: {text!r}"
        )
    if "structural" in block:
        structural = bool(block["structural"])
    else:
        structural = level == 1 and bool(re.match(
            r"^(?:СОДЕРЖАНИЕ|ВВЕДЕНИЕ|ЗАКЛЮЧЕНИЕ|ВЫВОДЫ|СПИСОК|ПРИЛОЖЕНИЕ)", text, re.I
        ))
    style = "Report Structural Heading" if structural else f"Heading {level}"
    paragraph = document.add_paragraph(text, style=style)
    if "page_break_before" in block:
        page_break_before = bool(block["page_break_before"])
    else:
        normalized_heading = re.sub(r"\s+", " ", text.strip().casefold())
        if level == 1 and normalized_heading in OPENING_PAGE_HEADINGS:
            page_break_before = False
        else:
            page_break_before = level == 1 and normalized_heading in DEFAULT_NEW_PAGE_HEADINGS
    paragraph.paragraph_format.page_break_before = page_break_before
    if structural:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.left_indent = Cm(0)
        paragraph.paragraph_format.first_line_indent = Cm(0)
    else:
        # Match the confirmed Word layout exactly. The direct first-line
        # position overrides application-specific list defaults, so the
        # number begins at 1.25 cm in both Word and LibreOffice.
        paragraph.paragraph_format.left_indent = Cm(0)
        paragraph.paragraph_format.first_line_indent = Cm(1.25)


def add_paragraph_block(document: Document, block: dict[str, Any]) -> None:
    paragraph = document.add_paragraph()
    alignments = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }
    alignment = str(block.get("align", "")).strip().lower()
    if alignment:
        if alignment not in alignments:
            raise SpecError(f"Unsupported paragraph alignment: {alignment}")
        if alignment == "left":
            raise SpecError(
                "Body prose must stay justified (по ширине); do not request left alignment. "
                "Rewrite or split paragraphs with long identifiers into shorter factual sentences instead."
            )
        paragraph.alignment = alignments[alignment]
    if "first_indent_cm" in block:
        paragraph.paragraph_format.first_line_indent = Cm(float(block["first_indent_cm"]))
    prefix = str(block.get("bold_prefix", ""))
    if PLACEHOLDER_PATTERN.search(prefix):
        raise SpecError(f"Placeholder detected in bold_prefix: {prefix[:80]}")
    if prefix and not prefix.endswith((" ", "\u00a0")):
        prefix += " "
    text = clean_text(block.get("text", ""), field="paragraph")
    if prefix:
        run = paragraph.add_run(prefix)
        set_run_font(run, "Times New Roman", 14, bold=True)
    if text:
        add_report_text(paragraph, text)
    keep_paragraph_lines(paragraph)


def add_list(document: Document, block: dict[str, Any], numbered: bool) -> None:
    style = "List Number" if numbered else "List Bullet"
    items = block.get("items")
    if not isinstance(items, list) or not items:
        raise SpecError(f"{block.get('type')} requires a non-empty items list")
    for item in items:
        paragraph = document.add_paragraph(style=style)
        add_report_text(paragraph, clean_text(item, field="list item"))
        keep_paragraph_lines(paragraph)


def add_figure(document: Document, block: dict[str, Any]) -> None:
    path = Path(clean_text(block.get("path"), field="figure path")).expanduser().resolve()
    if not path.is_file():
        raise SpecError(f"Figure does not exist: {path}")
    width = float(block.get("width_cm", 14.5))
    if width <= 0 or width > MAX_CONTENT_WIDTH_CM:
        raise SpecError(f"Figure width must be within 0..{MAX_CONTENT_WIDTH_CM} cm")
    keep_previous_reference(document)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    shape = run.add_picture(str(path), width=Cm(width))
    alt = str(block.get("alt_text", block.get("caption", ""))).strip()
    if alt:
        doc_pr = shape._inline.docPr
        doc_pr.set("descr", alt)
    number = clean_text(block.get("number", ""), field="figure number")
    if not number:
        raise SpecError("figure requires a number")
    caption = clean_text(block.get("caption", ""), field="figure caption")
    if not caption:
        raise SpecError("figure requires a caption")
    if caption.endswith("."):
        raise SpecError("figure caption must not end with a period")
    label = f"Рисунок {number} — {caption}"
    document.add_paragraph(label, style="Report Caption")


def set_paragraph_border(paragraph, *, size: int = 4, color: str = "000000") -> None:
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    for name in ("top", "left", "bottom", "right"):
        border = borders.find(qn(f"w:{name}"))
        if border is None:
            border = OxmlElement(f"w:{name}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), str(size))
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)


def add_figure_placeholder(document: Document, block: dict[str, Any]) -> None:
    number = clean_text(block.get("number", ""), field="figure placeholder number")
    caption = clean_text(block.get("caption", ""), field="figure placeholder caption")
    if not number or not caption:
        raise SpecError("figure_placeholder requires number and caption")
    width = float(block.get("width_cm", 15.5))
    height = float(block.get("height_cm", 6.0))
    if width <= 0 or width > MAX_CONTENT_WIDTH_CM:
        raise SpecError(f"Figure placeholder width must be within 0..{MAX_CONTENT_WIDTH_CM} cm")
    if height < 3.0 or height > 18.0:
        raise SpecError("Figure placeholder height must be within 3..18 cm")
    keep_previous_reference(document)
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    tbl_pr = table._tbl.tblPr
    descriptor = OxmlElement("w:tblCaption")
    descriptor.set(qn("w:val"), "figure-placeholder")
    tbl_pr.append(descriptor)
    row = table.rows[0]
    row.height = Cm(height)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    set_row_cant_split(row)
    cell = table.cell(0, 0)
    set_cell_width(cell, width)
    set_cell_margins(cell, top=0, start=100, bottom=0, end=100)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.style = document.styles["Report Figure Placeholder"]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    # Intentionally keep the reserved area empty. Editorial instructions such as
    # "insert a screenshot" must never leak into the submitted report.
    label = f"Рисунок {number} — {caption}"
    document.add_paragraph(label, style="Report Caption")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    existing = tr_pr.find(qn("w:tblHeader"))
    if existing is not None:
        tr_pr.remove(existing)
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_row_cant_split(row) -> None:
    """Keep a logical table row on one page."""
    tr_pr = row._tr.get_or_add_trPr()
    existing = tr_pr.find(qn("w:cantSplit"))
    if existing is None:
        existing = OxmlElement("w:cantSplit")
        tr_pr.append(existing)
    existing.set(qn("w:val"), "true")


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{tag}"))
        if element is None:
            element = OxmlElement(f"w:{tag}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(round(width_cm * 567)))


def set_cell_no_wrap(cell) -> None:
    """Keep an already width-checked exact token on one line."""
    tc_pr = cell._tc.get_or_add_tcPr()
    no_wrap = tc_pr.find(qn("w:noWrap"))
    if no_wrap is None:
        no_wrap = OxmlElement("w:noWrap")
        tc_pr.append(no_wrap)
    no_wrap.set(qn("w:val"), "true")


def choose_widths(headers: list[Any], rows: list[list[Any]]) -> list[float]:
    columns = len(headers)
    weights = []
    for index in range(columns):
        values = [str(headers[index])] + [str(row[index]) if index < len(row) else "" for row in rows]
        weights.append(max(4, min(40, max(map(len, values), default=4))))
    total = sum(weights)
    return [MAX_CONTENT_WIDTH_CM * weight / total for weight in weights]


def add_table_part(
    document: Document,
    headers: list[Any],
    rows: list[list[Any]],
    widths: list[float],
    alignments: list[str] | None = None,
    font_size: float = 14,
    font_family: str = "Times New Roman",
    line_spacing: float = 1.5,
    no_wrap_columns: set[int] | None = None,
    keep_together: bool = False,
):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    table_width = sum(map(float, widths))
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(round(table_width * 567)))

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(round(float(width) * 567)))
        grid.append(col)

    all_rows = [headers] + rows
    for row_index, values in enumerate(all_rows):
        row = table.rows[0] if row_index == 0 else table.add_row()
        set_row_cant_split(row)
        if row_index == 0:
            set_repeat_table_header(row)
        for index, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            compact_cell = row_index > 0 and bool(no_wrap_columns) and index in no_wrap_columns
            set_cell_width(cell, float(width))
            if compact_cell:
                set_cell_margins(cell, top=40, start=40, bottom=40, end=40)
            else:
                set_cell_margins(cell)
            if no_wrap_columns and index in no_wrap_columns:
                set_cell_no_wrap(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            if row_index == 0:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.keep_with_next = True
            elif alignments:
                value_text = str(value).strip().replace("−", "-")
                if re.fullmatch(r"[+-]?\d+(?:[.,]\d+)?", value_text):
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    paragraph.alignment = {
                        "left": WD_ALIGN_PARAGRAPH.LEFT,
                        "center": WD_ALIGN_PARAGRAPH.CENTER,
                        "right": WD_ALIGN_PARAGRAPH.RIGHT,
                    }[alignments[index]]
            else:
                value_text = str(value).strip().replace("−", "-")
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                    if re.fullmatch(r"[+-]?\d+(?:[.,]\d+)?", value_text)
                    else WD_ALIGN_PARAGRAPH.LEFT
                )
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing = line_spacing if compact_cell else 1.5
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(clean_text(value, field=f"table cell {row_index},{index}"))
            run_font = font_family if compact_cell else "Times New Roman"
            run_size = font_size if compact_cell else 14
            set_run_font(run, run_font, run_size, bold=(row_index == 0))

    if keep_together and len(table.rows) > 1:
        # Word and LibreOffice honour keep-with-next inside table cells. Applying
        # it to every row except the last keeps a short logical table intact;
        # the option must not be used for a table taller than one page.
        for row in table.rows[:-1]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True

    return table


def add_table(document: Document, block: dict[str, Any]) -> None:
    headers = block.get("headers")
    rows = block.get("rows")
    if not isinstance(headers, list) or not headers:
        raise SpecError("table requires non-empty headers")
    if not isinstance(rows, list):
        raise SpecError("table rows must be a list")
    if any(not isinstance(row, list) or len(row) != len(headers) for row in rows):
        raise SpecError("every table row must have the same number of cells as headers")
    widths = [float(value) for value in (block.get("widths_cm") or choose_widths(headers, rows))]
    if len(widths) != len(headers) or sum(widths) > MAX_CONTENT_WIDTH_CM + 0.01:
        raise SpecError(f"table widths must match columns and total at most {MAX_CONTENT_WIDTH_CM} cm")
    for header, width in zip(headers, widths):
        if str(header).strip() == "Итер." and width < 1.8:
            raise SpecError("the `Итер.` column must be at least 1.8 cm at Times New Roman 14 pt")
    alignments = block.get("column_alignments")
    if alignments is not None:
        if (
            not isinstance(alignments, list)
            or len(alignments) != len(headers)
            or any(value not in {"left", "center", "right"} for value in alignments)
        ):
            raise SpecError("column_alignments must match columns and use left, center, or right")
    long_token_mode = bool(block.get("long_token_mode", False))
    font_size = float(block.get("font_size", 14))
    font_family = str(block.get("font_family", "Times New Roman"))
    line_spacing = float(block.get("line_spacing", 1.5))
    if font_size != 14.0 or font_family != "Times New Roman" or line_spacing != 1.5:
        raise SpecError(
            "report tables must use Times New Roman 14 pt with 1.5 line spacing; "
            "fit long tokens by column structure and natural Word wrapping in portrait pages"
        )
    no_wrap_raw = block.get("no_wrap_columns", [])
    if not isinstance(no_wrap_raw, list) or any(not isinstance(value, int) for value in no_wrap_raw):
        raise SpecError("no_wrap_columns must be a list of zero-based integer column indexes")
    no_wrap_columns = set(no_wrap_raw)
    if any(value < 0 or value >= len(headers) for value in no_wrap_columns):
        raise SpecError("no_wrap_columns contains an index outside the table")
    if no_wrap_columns or long_token_mode:
        raise SpecError(
            "long_token_mode and no_wrap_columns are disabled: preserve portrait pages and "
            "allow exact values to wrap naturally inside Word table cells"
        )

    number = clean_text(block.get("number", ""), field="table number")
    if not number:
        raise SpecError("table requires a number")
    caption = clean_text(block.get("caption", ""), field="table caption")
    if not caption:
        raise SpecError("table requires a caption")
    if caption.endswith("."):
        raise SpecError("table caption must not end with a period")

    rows_per_page_raw = block.get("rows_per_page")
    if rows_per_page_raw not in (None, ""):
        raise SpecError(
            "rows_per_page is disabled: keep one native Word table and use repeated headers"
        )
    parts = [rows]

    keep_together_raw = block.get("keep_together")
    if bool(keep_together_raw):
        raise SpecError(
            "keep_together is disabled for whole tables because it creates large blank areas"
        )
    keep_together = False

    use_final_label = bool(block.get("use_final_label", False))
    if use_final_label:
        raise SpecError("use_final_label is disabled for native Word table pagination")

    keep_previous_reference(document)
    for part_index, part_rows in enumerate(parts):
        if part_index == 0:
            label = f"Таблица {number} — {caption}"
        elif use_final_label and part_index == len(parts) - 1:
            label = f"Окончание таблицы {number}"
        else:
            label = f"Продолжение таблицы {number}"
        caption_paragraph = document.add_paragraph(label, style="Report Table Caption")
        caption_paragraph.paragraph_format.space_before = Pt(0)
        caption_paragraph.paragraph_format.space_after = Pt(0)
        add_table_part(
            document,
            headers,
            part_rows,
            widths,
            alignments,
            font_size,
            font_family,
            line_spacing,
            no_wrap_columns,
            keep_together=keep_together and len(parts) == 1,
        )
        if part_index < len(parts) - 1:
            document.add_page_break()



@functools.lru_cache(maxsize=256)
def latex_to_omath(latex: str):
    pandoc = shutil.which("pandoc")
    if not pandoc:
        raise SpecError("Pandoc is required to create native Microsoft Word equations")
    source = f"$$\n{latex}\n$$\n"
    with tempfile.TemporaryDirectory(prefix="report-equation-") as tmp:
        output = Path(tmp) / "equation.docx"
        process = subprocess.run(
            [pandoc, "--from", "markdown", "--to", "docx", "--output", str(output)],
            input=source,
            text=True,
            capture_output=True,
        )
        if process.returncode != 0 or not output.is_file():
            message = process.stderr.strip() or "unknown Pandoc error"
            raise SpecError(f"Cannot convert LaTeX equation to OMML: {message}")
        with zipfile.ZipFile(output) as archive:
            xml = archive.read("word/document.xml")
    root = etree.fromstring(xml)
    namespaces = {"m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}
    nodes = root.xpath(".//m:oMath", namespaces=namespaces)
    if not nodes:
        raise SpecError("Pandoc produced no OMML equation")
    return copy.deepcopy(nodes[0])


def add_equation(document: Document, block: dict[str, Any]) -> None:
    paragraph = document.add_paragraph(style="Report Equation")
    number = str(block.get("number", "")).strip()
    if number:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.tab_stops.add_tab_stop(
            Cm(MAX_CONTENT_WIDTH_CM / 2), WD_TAB_ALIGNMENT.CENTER
        )
        paragraph.paragraph_format.tab_stops.add_tab_stop(
            Cm(MAX_CONTENT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT
        )
        paragraph.add_run().add_tab()
    latex = clean_text(block.get("latex"), field="equation latex")
    if re.search(r"\\\\|\\begin\{(?:aligned|alignedat|array|matrix|pmatrix|bmatrix|vmatrix|Vmatrix)\}", latex):
        raise SpecError(
            "one equation block must contain one equation or parameter value; "
            "create separate equation blocks instead of aligned arrays or matrices"
        )
    longest_number = max((len(value) for value in re.findall(r"\d+", latex)), default=0)
    if longest_number > 60:
        raise SpecError(
            "an exact integer longer than 60 digits must be represented by named decimal "
            "blocks in separate 14 pt equations"
        )
    equation_half_points = "28"
    paragraph.style = document.styles[f"Report Equation {equation_half_points}"]
    equation = copy.deepcopy(latex_to_omath(latex))
    for math_run in equation.xpath(".//m:r"):
        rpr = math_run.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            math_run.insert(0, rpr)
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{key}"), "Cambria Math")
        for tag in ("w:sz", "w:szCs"):
            size = rpr.find(qn(tag))
            if size is None:
                size = OxmlElement(tag)
                rpr.append(size)
            size.set(qn("w:val"), equation_half_points)
        character_spacing = rpr.find(qn("w:spacing"))
        if character_spacing is not None:
            rpr.remove(character_spacing)
    paragraph_run_properties = paragraph._p.get_or_add_pPr().find(qn("w:rPr"))
    if paragraph_run_properties is None:
        paragraph_run_properties = OxmlElement("w:rPr")
        paragraph._p.get_or_add_pPr().append(paragraph_run_properties)
    paragraph_fonts = paragraph_run_properties.find(qn("w:rFonts"))
    if paragraph_fonts is None:
        paragraph_fonts = OxmlElement("w:rFonts")
        paragraph_run_properties.insert(0, paragraph_fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        paragraph_fonts.set(qn(f"w:{key}"), "Cambria Math")
    for tag in ("w:sz", "w:szCs"):
        size = paragraph_run_properties.find(qn(tag))
        if size is None:
            size = OxmlElement(tag)
            paragraph_run_properties.append(size)
        size.set(qn("w:val"), equation_half_points)
    paragraph._p.append(equation)
    if number:
        run = paragraph.add_run()
        run.add_tab()
        run.add_text(f"({number})")
        set_run_font(run, "Times New Roman", 14)


def add_code(document: Document, block: dict[str, Any]) -> None:
    source = str(block.get("source", "")).strip()
    if source:
        path = Path(source).expanduser().resolve()
        if not path.is_file():
            raise SpecError(f"Code source does not exist: {path}")
        text = path.read_text(encoding=str(block.get("encoding", "utf-8")), errors="replace")
    else:
        text = str(block.get("text", ""))
    if PLACEHOLDER_PATTERN.search(text):
        raise SpecError("Placeholder detected in code block")
    presentation = str(block.get("presentation", "fragment")).strip().lower()
    if presentation not in {"fragment", "appendix"}:
        raise SpecError("code presentation must be 'fragment' or 'appendix'")
    caption = str(block.get("caption", "")).strip()
    number = str(block.get("number", "")).strip()
    if presentation == "fragment":
        if not caption or not number:
            raise SpecError("framed code fragment requires caption and number")
        if caption.endswith("."):
            raise SpecError("listing caption must not end with a period")
        keep_previous_reference(document)
        paragraph = document.add_paragraph(f"Листинг {number} — {caption}", style="Report Table Caption")
        paragraph.paragraph_format.keep_with_next = True
    lines = text.expandtabs(4).splitlines() or [""]
    if presentation == "fragment":
        font_size = float(block.get("font_size", 12))
        if font_size not in {11.0, 12.0}:
            raise SpecError("framed code fragment font_size must be 11 or 12 pt")
        table = document.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.style = "Table Grid"
        tbl_pr = table._tbl.tblPr
        descriptor = OxmlElement("w:tblCaption")
        descriptor.set(qn("w:val"), "code-fragment")
        tbl_pr.append(descriptor)
        set_row_cant_split(table.rows[0])
        cell = table.cell(0, 0)
        set_cell_width(cell, MAX_CONTENT_WIDTH_CM)
        set_cell_margins(cell, top=70, start=100, bottom=70, end=100)
        for index, line in enumerate(lines):
            paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
            paragraph.style = document.styles["Report Code Fragment"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            set_run_font(paragraph.add_run(line), "Courier New", font_size)
        document.add_paragraph().paragraph_format.space_after = Pt(0)
    else:
        label = clean_text(block.get("label", ""), field="appendix code label")
        if not label:
            raise SpecError("appendix code requires a visible relative-path label")
        if Path(label).is_absolute() or ".." in Path(label).parts:
            raise SpecError("appendix code label must be a safe relative path")
        label_paragraph = document.add_paragraph(style="Report Appendix File Label")
        label_run = label_paragraph.add_run(f"Файл: {label}")
        set_run_font(label_run, "Times New Roman", 12, bold=True)
        paragraph = document.add_paragraph(style="Report Appendix Code")
        line_spacing_pt = block.get("line_spacing_pt")
        if line_spacing_pt is None:
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            paragraph.paragraph_format.line_spacing = 1.0
        else:
            line_spacing_pt = float(line_spacing_pt)
            if line_spacing_pt < 10.0 or line_spacing_pt > 12.0:
                raise SpecError("appendix code line_spacing_pt must be within 10..12 pt")
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            paragraph.paragraph_format.line_spacing = Pt(line_spacing_pt)
        paragraph.paragraph_format.widow_control = True
        run = paragraph.add_run()
        for index, line in enumerate(lines):
            if index:
                run.add_break()
            run.add_text(line)
        set_run_font(run, "Courier New", 10)


def add_appendix_heading(document: Document, block: dict[str, Any]) -> None:
    letter = clean_text(block.get("letter"), field="appendix letter").upper()
    title = clean_text(block.get("title", ""), field="appendix title")
    paragraph = document.add_paragraph(f"Приложение {letter}", style="Report Structural Heading")
    paragraph.paragraph_format.page_break_before = True
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    if title:
        # Inherit the justified body alignment; a short single-line title
        # still renders flush left because it is the paragraph's last line.
        document.add_paragraph(title, style="Normal")


def add_bibliography(document: Document, block: dict[str, Any]) -> None:
    title = clean_text(block.get("title", "Список использованных источников"), field="bibliography title")
    heading = document.add_paragraph(title, style="Report Structural Heading")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.first_line_indent = Cm(0)
    items = block.get("items")
    if not isinstance(items, list) or not items:
        raise SpecError("bibliography requires at least one verified source")
    for index, item in enumerate(items, 1):
        if not isinstance(item, dict):
            raise SpecError(
                "Every bibliography item must be an object with kind, citation, url, accessed, and verified=true"
            )
        kind = str(item.get("kind", "")).strip()
        if kind not in SOURCE_KINDS:
            raise SpecError(f"Unsupported bibliography source kind: {kind!r}")
        if item.get("verified") is not True:
            raise SpecError("Every bibliography source must be opened and verified before assembly")
        citation = clean_text(item.get("citation"), field="bibliography citation").rstrip(". ")
        url = clean_text(item.get("url"), field="bibliography url").rstrip(".,;:)")
        parsed = urlparse(url)
        host = (parsed.hostname or "").casefold()
        if parsed.scheme != "https" or not host or host in {"localhost", "127.0.0.1", "::1"}:
            raise SpecError(f"Bibliography URL must be a public HTTPS link: {url!r}")
        accessed = clean_text(item.get("accessed"), field="bibliography access date")
        try:
            dt.datetime.strptime(accessed, "%d.%m.%Y")
        except ValueError as exc:
            raise SpecError(f"Bibliography accessed date must use DD.MM.YYYY: {accessed!r}") from exc
        paragraph = document.add_paragraph(style="Report Bibliography")
        bookmark_start = OxmlElement("w:bookmarkStart")
        bookmark_start.set(qn("w:id"), str(50000 + index))
        bookmark_start.set(qn("w:name"), f"_ReportBib{index}")
        paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, bookmark_start)
        prefix = paragraph.add_run(f"{index}. {citation}. URL: ")
        set_run_font(prefix, "Times New Roman", 14)
        add_external_hyperlink(paragraph, url, url)
        suffix = paragraph.add_run(f" (дата обращения: {accessed}).")
        set_run_font(suffix, "Times New Roman", 14)
        bookmark_end = OxmlElement("w:bookmarkEnd")
        bookmark_end.set(qn("w:id"), str(50000 + index))
        paragraph._p.append(bookmark_end)


def add_content(document: Document, blocks: Iterable[dict[str, Any]]) -> None:
    handlers = {
        "heading": add_heading,
        "paragraph": add_paragraph_block,
        "bullet_list": lambda doc, block: add_list(doc, block, False),
        "numbered_list": lambda doc, block: add_list(doc, block, True),
        "figure": add_figure,
        "figure_placeholder": add_figure_placeholder,
        "table": add_table,
        "equation": add_equation,
        "code": add_code,
        "appendix_heading": add_appendix_heading,
        "bibliography": add_bibliography,
    }
    for index, block in enumerate(blocks):
        if not isinstance(block, dict):
            raise SpecError(f"content[{index}] must be an object")
        block_type = block.get("type")
        if block_type == "page_break":
            document.add_page_break()
            continue
        handler = handlers.get(block_type)
        if handler is None:
            raise SpecError(f"Unsupported content block type at index {index}: {block_type}")
        handler(document, block)


def validate_spec(spec: dict[str, Any]) -> None:
    if spec.get("report_type") not in {"lab", "course"}:
        raise SpecError("report_type must be 'lab' or 'course'")
    if not isinstance(spec.get("metadata"), dict):
        raise SpecError("metadata must be an object")
    if not isinstance(spec.get("content"), list) or not spec["content"]:
        raise SpecError("content must be a non-empty array")
    document_settings = spec.get("document") or {}
    if not isinstance(document_settings, dict):
        raise SpecError("document must be an object")
    code_change = document_settings.get("code_change")
    if code_change is not None:
        if not isinstance(code_change, dict):
            raise SpecError("document.code_change must be an object")
        mode = str(code_change.get("mode", "")).strip().lower()
        if mode not in {"new", "rewrite", "refactor", "modify", "none"}:
            raise SpecError("document.code_change.mode must be new, rewrite, refactor, modify, or none")
        source_keys = ("original_sources", "final_sources", "test_sources", "supporting_sources")
        declared_sources: list[Path] = []
        groups: dict[str, list[Path]] = {}
        for source_key in source_keys:
            raw_sources = code_change.get(source_key, [])
            if not isinstance(raw_sources, list):
                raise SpecError(f"document.code_change.{source_key} must be an array")
            resolved_sources: list[Path] = []
            for raw_source in raw_sources:
                source_path = Path(str(raw_source)).expanduser().resolve()
                if not source_path.is_file():
                    raise SpecError(f"Declared code source does not exist: {source_path}")
                resolved_sources.append(source_path)
                declared_sources.append(source_path)
            groups[source_key] = resolved_sources
        if mode in {"new", "rewrite", "modify"} and not groups["final_sources"]:
            raise SpecError(f"document.code_change mode {mode!r} requires final_sources")
        if mode == "refactor" and (not groups["original_sources"] or not groups["final_sources"]):
            raise SpecError("document.code_change mode 'refactor' requires original_sources and final_sources")
        if mode == "none" and declared_sources:
            raise SpecError("document.code_change mode 'none' cannot declare source files")
        if len(set(declared_sources)) != len(declared_sources):
            raise SpecError("document.code_change lists contain a duplicate source path")
        appendix_sources: list[Path] = []
        appendix_labels: list[str] = []
        for block in spec["content"]:
            if not isinstance(block, dict) or block.get("type") != "code":
                continue
            if str(block.get("presentation", "fragment")).strip().lower() != "appendix":
                continue
            source_value = str(block.get("source", "")).strip()
            if source_value:
                appendix_sources.append(Path(source_value).expanduser().resolve())
            appendix_labels.append(str(block.get("label", "")).strip())
        if mode != "none" and not any(
            isinstance(block, dict) and block.get("type") == "appendix_heading"
            for block in spec["content"]
        ):
            raise SpecError("programming report requires at least one appendix_heading")
        for source_path in declared_sources:
            if appendix_sources.count(source_path) != 1:
                raise SpecError(
                    f"Declared code source must appear exactly once as appendix code: {source_path}"
                )
        if mode != "none" and any(not label for label in appendix_labels):
            raise SpecError("every appendix code block requires a visible relative-path label")
    report_type = spec["report_type"]
    bibliography_blocks = [
        block for block in spec["content"]
        if isinstance(block, dict) and block.get("type") == "bibliography"
    ]
    if report_type == "course" and len(bibliography_blocks) != 1:
        raise SpecError("course content must contain exactly one bibliography with verified real sources")
    if report_type == "lab" and bibliography_blocks:
        raise SpecError("laboratory reports must not contain a bibliography or external source hyperlinks")
    bibliography_count = len(bibliography_blocks[0].get("items", [])) if bibliography_blocks else 0
    citation_numbers: list[int] = []
    for block in spec["content"]:
        if not isinstance(block, dict):
            continue
        values: list[str] = []
        if block.get("type") == "paragraph":
            values.append(str(block.get("text", "")))
        elif block.get("type") in {"bullet_list", "numbered_list"}:
            values.extend(str(item) for item in block.get("items", []))
        for value in values:
            citation_numbers.extend(int(number) for number in CITATION_PATTERN.findall(value))
    if report_type == "lab" and citation_numbers:
        raise SpecError("laboratory reports must not contain numeric source citations such as [1]")
    if report_type == "course" and any(number < 1 or number > bibliography_count for number in citation_numbers):
        raise SpecError("course-paper citation [N] refers to a missing bibliography item")
    previous_numbered_level = 0
    for index, block in enumerate(spec["content"]):
        if not isinstance(block, dict) or block.get("type") != "heading":
            continue
        level = int(block.get("level", 1))
        text = clean_text(block.get("text"), field=f"content[{index}] heading")
        if MANUAL_HEADING_NUMBER.match(text):
            raise SpecError(
                f"content[{index}] heading contains a manual number; supply only the title text"
            )
        explicit_structural = block.get("structural") if "structural" in block else None
        inferred_structural = level == 1 and bool(re.match(
            r"^(?:СОДЕРЖАНИЕ|ВВЕДЕНИЕ|ЗАКЛЮЧЕНИЕ|ВЫВОДЫ|СПИСОК|ПРИЛОЖЕНИЕ)", text, re.I
        ))
        structural = bool(explicit_structural) if explicit_structural is not None else inferred_structural
        if structural:
            continue
        if level > previous_numbered_level + 1:
            raise SpecError(f"Heading hierarchy skips a level at content[{index}]: {text!r}")
        previous_numbered_level = level
    for key in (
        "title", "discipline", "student_group", "student_name",
        "reviewer_position", "reviewer_name", "city", "year",
    ):
        required(spec["metadata"], key)


def load_json_object(path: Path, label: str) -> dict[str, Any]:
    path = path.expanduser().resolve()
    if not path.is_file():
        raise SpecError(f"{label} does not exist: {path}")
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise SpecError(f"{label} is not valid JSON: {exc}") from exc
    if not isinstance(value, dict):
        raise SpecError(f"{label} must be a JSON object")
    return value


def normalize_discipline(value: Any) -> str:
    return str(value or "").strip().strip("«»\"").casefold()


def merge_metadata_sources(
    spec: dict[str, Any],
    profile_path: Path | None,
    subjects_path: Path | None,
) -> dict[str, Any]:
    current = spec.get("metadata")
    if not isinstance(current, dict):
        raise SpecError("metadata must be an object")
    merged: dict[str, Any] = {}

    if profile_path is not None:
        profile = load_json_object(profile_path, "Metadata profile")
        forbidden = sorted(key for key in profile if key.startswith("reviewer_"))
        if forbidden:
            raise SpecError(
                "Global metadata profile must not contain discipline-specific reviewer fields: "
                + ", ".join(forbidden)
            )
        merged.update(profile)

    discipline = current.get("discipline") or merged.get("discipline")
    if subjects_path is not None:
        subjects = load_json_object(subjects_path, "Subject profile")
        normalized = normalize_discipline(discipline)
        if not normalized:
            raise SpecError("discipline is required to select a subject-specific reviewer")
        matches = [value for key, value in subjects.items() if normalize_discipline(key) == normalized]
        if not matches:
            raise SpecError(f"No reviewer profile found for discipline: {discipline}")
        subject = matches[0]
        if not isinstance(subject, dict):
            raise SpecError(f"Reviewer profile for discipline {discipline!r} must be an object")
        allowed = {"reviewer_role", "reviewer_position", "reviewer_name"}
        unexpected = sorted(set(subject) - allowed)
        if unexpected:
            raise SpecError(
                f"Reviewer profile for discipline {discipline!r} has unsupported fields: "
                + ", ".join(unexpected)
            )
        merged.update(subject)

    for key, value in current.items():
        if value not in (None, ""):
            merged[key] = value
    result = dict(spec)
    result["metadata"] = merged
    return result


def build(spec: dict[str, Any], output: Path) -> None:
    validate_spec(spec)
    document = Document()
    settings = spec.get("document") or {}
    configure_document(document, int(settings.get("page_number_start", 0)))

    properties = document.core_properties
    properties.title = required(spec["metadata"], "title")
    properties.subject = str(settings.get("subject", "Учебный отчёт"))
    properties.author = str(settings.get("author", ""))
    properties.last_modified_by = ""
    properties.keywords = ""
    properties.comments = ""

    add_cover(document, spec)
    include_toc = bool(settings.get("include_toc", True))
    if include_toc:
        levels = max(1, min(3, int(settings.get("toc_levels", 3))))
        add_toc(document, levels)
    content = [dict(block) for block in spec["content"]]
    if include_toc and content:
        if content[0].get("type") == "heading":
            content[0]["page_break_before"] = True
        else:
            content.insert(0, {"type": "page_break"})
    add_content(document, content)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--spec", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--profile", type=Path, help="Optional reusable .report-profile.json")
    parser.add_argument("--subjects", type=Path, help="Optional discipline reviewers .report-subjects.json")
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()

    spec_path = args.spec.expanduser().resolve()
    output = args.output.expanduser().resolve()
    if output.suffix.lower() != ".docx":
        parser.error("The final output must use the .docx extension")
    if output.exists() and not args.force:
        parser.error(f"Output already exists; pass --force to replace it: {output}")
    try:
        spec = json.loads(spec_path.read_text(encoding="utf-8"))
        built_in_assets = Path(__file__).resolve().parent.parent / "assets"
        profile_path = args.profile
        if profile_path is None and isinstance(spec, dict) and spec.get("profile_path"):
            profile_path = Path(str(spec["profile_path"]))
        if profile_path is None:
            candidate = built_in_assets / "report-profile.default.json"
            if candidate.is_file():
                profile_path = candidate
        subjects_path = args.subjects
        if subjects_path is None and isinstance(spec, dict) and spec.get("subject_profile_path"):
            subjects_path = Path(str(spec["subject_profile_path"]))
        current_metadata = spec.get("metadata") if isinstance(spec, dict) else None
        reviewer_is_missing = not isinstance(current_metadata, dict) or any(
            not str(current_metadata.get(key, "")).strip()
            for key in ("reviewer_position", "reviewer_name")
        )
        if subjects_path is None and reviewer_is_missing:
            candidate = built_in_assets / "report-subjects.default.json"
            if candidate.is_file():
                subjects_path = candidate
        spec = merge_metadata_sources(spec, profile_path, subjects_path)
        build(spec, output)
    except (OSError, json.JSONDecodeError, SpecError) as exc:
        parser.error(str(exc))
    print(json.dumps({"output": str(output), "format": "docx"}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
